#!/usr/bin/env python3
"""Build certified annual L5 voluntary termination-rate facts for 2022-2024."""
from __future__ import annotations

import hashlib
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve()
HKIA_ROOT = HERE.parents[3]
SOURCE_ROOT = HKIA_ROOT / "12_分析框架验证_Validate_Framework/01_sources/raw/SRC-REG-IA-LTA"
OUTPUT = HERE.parents[1] / "data/annual_market_fact_layer_2022_2024.db"

PRODUCTS = {"whole life": "whole_life", "whole of life": "whole_life",
            "endowment": "endowment", "all policies": "all_policies"}


def product_semantic(value) -> str | None:
    """Map bilingual labels without relying on newline/slash placement."""
    text = " ".join(str(value or "").replace("/", " ").replace("*", " ").split()).lower()
    for label, semantic in PRODUCTS.items():
        if label in text:
            return semantic
    return None


def status_value(value):
    if value is None or str(value).strip() == "":
        return None, "blank"
    if isinstance(value, str) and ("N.A" in value.upper() or "不適用" in value):
        return None, "not_applicable"
    if isinstance(value, (int, float)):
        return float(value), "reported_zero" if float(value) == 0 else "reported"
    return None, "unparsed"


def add(rows, *, year, linked, product, metric, observation_year, policy_year_band,
        participation, value, cell, source_file, checksum):
    numeric, record_status = status_value(value)
    rows.append((
        f"L5:{year}:{linked}:{product}:{metric}:{observation_year or year}:"
        f"{policy_year_band or 'all'}:{participation or 'all'}",
        year, observation_year, "L5", linked, product, metric,
        policy_year_band, participation, numeric, "percentage_point", record_status,
        "certified", "annual_lt", source_file.name, source_file.name, cell, checksum,
    ))


def parse_year(year: int):
    folder = SOURCE_ROOT / str(year) / "full_annual_set"
    source_file = next(folder.glob(f"Table-L5*{year}.xlsx"))
    checksum = hashlib.sha256(source_file.read_bytes()).hexdigest()
    ws = load_workbook(source_file, data_only=True, read_only=False).active
    rows = []

    # 2022/2023 use rows 10/11/13 and 21/22/24; 2024 inserts one header row.
    non_rows = (11, 12, 14) if year == 2024 else (10, 11, 13)
    linked_rows = (22, 23, 25) if year == 2024 else (21, 22, 24)
    non_year_row = 9 if year == 2024 else 8
    linked_year_row = 20 if year == 2024 else 19

    for linked, data_rows, year_row in (
        ("non_linked", non_rows, non_year_row),
        ("linked", linked_rows, linked_year_row),
    ):
        for row_number in data_rows:
            product_label = ws.cell(row_number, 2).value
            product = product_semantic(product_label)
            if product is None:
                raise ValueError(f"Unmapped product {product_label!r} in {source_file}:{row_number}")

            # Five-year overall series.
            for col in range(3, 8):
                observation_year = ws.cell(year_row, col).value
                if not isinstance(observation_year, int):
                    raise ValueError(f"Invalid observation year {source_file}:{ws.cell(year_row,col).coordinate}")
                cell = ws.cell(row_number, col)
                add(rows, year=year, linked=linked, product=product,
                    metric="overall_termination_rate", observation_year=observation_year,
                    policy_year_band=None, participation=None, value=cell.value,
                    cell=cell.coordinate, source_file=source_file, checksum=checksum)

            # Latest report-year policy-duration breakdown. Non-linked splits participating status.
            if linked == "non_linked":
                for band, cols in (("policy_year_1", (8, 9)), ("policy_year_2", (11, 12)),
                                   ("policy_year_3_plus", (14, 15))):
                    for participation, col in zip(("participating", "other"), cols):
                        cell = ws.cell(row_number, col)
                        add(rows, year=year, linked=linked, product=product,
                            metric="termination_rate", observation_year=year,
                            policy_year_band=band, participation=participation, value=cell.value,
                            cell=cell.coordinate, source_file=source_file, checksum=checksum)
            else:
                for band, col in (("policy_year_1", 8), ("policy_year_2", 11),
                                  ("policy_year_3_plus", 14)):
                    cell = ws.cell(row_number, col)
                    add(rows, year=year, linked=linked, product=product,
                        metric="termination_rate", observation_year=year,
                        policy_year_band=band, participation="not_applicable", value=cell.value,
                        cell=cell.coordinate, source_file=source_file, checksum=checksum)
    return rows


def main():
    all_rows = []
    for year in (2022, 2023, 2024):
        all_rows.extend(parse_year(year))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        OUTPUT.unlink()
    con = sqlite3.connect(OUTPUT)
    con.executescript("""
    CREATE TABLE termination_rate_facts (
      fact_id TEXT PRIMARY KEY,
      report_year INTEGER NOT NULL,
      observation_year INTEGER NOT NULL,
      table_id TEXT NOT NULL,
      linked_status TEXT NOT NULL,
      insurance_type TEXT NOT NULL,
      metric_id TEXT NOT NULL,
      policy_year_band TEXT,
      participation_status TEXT,
      value REAL,
      unit TEXT NOT NULL,
      record_status TEXT NOT NULL,
      certification TEXT NOT NULL,
      schema_version TEXT NOT NULL,
      source_asset_id TEXT NOT NULL,
      source_file TEXT NOT NULL,
      source_locator TEXT NOT NULL,
      checksum_sha256 TEXT NOT NULL
    );
    CREATE INDEX idx_l5_period ON termination_rate_facts(observation_year, linked_status, insurance_type);
    CREATE INDEX idx_l5_report ON termination_rate_facts(report_year, metric_id);
    """)
    con.executemany("INSERT INTO termination_rate_facts VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", all_rows)
    con.commit()
    counts = con.execute("SELECT report_year, COUNT(*) FROM termination_rate_facts GROUP BY report_year").fetchall()
    statuses = con.execute("SELECT record_status, COUNT(*) FROM termination_rate_facts GROUP BY record_status").fetchall()
    con.close()
    print(f"built={OUTPUT}")
    print(f"rows={len(all_rows)} years={counts} statuses={statuses}")


if __name__ == "__main__":
    main()
