#!/usr/bin/env python3
"""Build certified L1 annual market facts for 2022-2024."""
from __future__ import annotations

import hashlib
import re
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve()
HKIA_ROOT = HERE.parents[3]
SOURCE_ROOT = HKIA_ROOT / "12_分析框架验证_Validate_Framework/01_sources/raw/SRC-REG-IA-LTA"
OUTPUT = HERE.parents[1] / "data/annual_market_fact_layer_2022_2024.db"

PRODUCT_ROWS = {11: "whole_life", 12: "endowment", 13: "term", 14: "other"}
PRODUCT_ROWS_2 = {25: "whole_life", 26: "endowment", 27: "term", 28: "other"}
PRODUCT_ROWS_NB = {47: "whole_life", 48: "endowment", 49: "term", 50: "other"}


def source_file(year: int) -> Path:
    folder = SOURCE_ROOT / str(year) / "full_annual_set"
    return next(p for p in folder.glob("Table-L*.xlsx") if re.match(r"Table-L1(?:_|-)", p.name))


def status_value(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "blank"
    if isinstance(value, str):
        normalized = value.upper()
        if "N.A" in normalized or "不適用" in value:
            return None, "not_applicable"
        if value.strip() == "-":
            return 0.0, "reported_zero"
        return None, "unparsed"
    if isinstance(value, (int, float)):
        return float(value), "reported_zero" if float(value) == 0 else "reported"
    return None, "unparsed"


def emit(rows, ws, *, report_year, observation_year, section, metric, unit,
         linked_status, insurance_type, component_type, row, col, source, checksum):
    value, status = status_value(ws.cell(row, col).value)
    locator = ws.cell(row, col).coordinate
    fact_id = ":".join(map(str, ("L1", report_year, observation_year, section, metric,
                                 linked_status, insurance_type, component_type, locator)))
    rows.append((fact_id, report_year, observation_year, "L1", section, metric, unit,
                 linked_status, insurance_type, component_type, value, status, "certified",
                 "rbc" if report_year >= 2024 else "pre_rbc", source.name, source.name,
                 ws.title, locator, checksum))


def metric_block(rows, ws, *, report_year, section, metric, unit, row_map, year_row,
                 start_col, linked_status, component_type, source, checksum):
    for row, insurance_type in row_map.items():
        for offset in range(5):
            col = start_col + offset
            observation_year = ws.cell(year_row, col).value
            if not isinstance(observation_year, int):
                raise ValueError(f"Invalid observation year at {source}:{ws.cell(year_row,col).coordinate}")
            emit(rows, ws, report_year=report_year, observation_year=observation_year,
                 section=section, metric=metric, unit=unit, linked_status=linked_status,
                 insurance_type=insurance_type, component_type=component_type, row=row, col=col,
                 source=source, checksum=checksum)


def parse_year(year: int):
    source = source_file(year)
    checksum = hashlib.sha256(source.read_bytes()).hexdigest()
    ws = load_workbook(source, data_only=True, read_only=False).active
    rows = []

    # L1a: policy count and office/revenue premium.
    for metric, unit, start_col in (("policy_count", "count", 4),
                                    ("office_or_revenue_premium", "HKD_million", 9)):
        metric_block(rows, ws, report_year=year, section="inforce", metric=metric, unit=unit,
                     row_map=PRODUCT_ROWS, year_row=8, start_col=start_col,
                     linked_status="non_linked", component_type="product", source=source, checksum=checksum)
        metric_block(rows, ws, report_year=year, section="inforce", metric=metric, unit=unit,
                     row_map={15: "all_products"}, year_row=8, start_col=start_col,
                     linked_status="non_linked", component_type="subtotal", source=source, checksum=checksum)
        metric_block(rows, ws, report_year=year, section="inforce", metric=metric, unit=unit,
                     row_map={17: "all_products"}, year_row=8, start_col=start_col,
                     linked_status="linked", component_type="subtotal", source=source, checksum=checksum)
        metric_block(rows, ws, report_year=year, section="inforce", metric=metric, unit=unit,
                     row_map={19: "all_products"}, year_row=8, start_col=start_col,
                     linked_status="total", component_type="market_total", source=source, checksum=checksum)

    # L1a: sums assured and liability/current estimate, including adjustment rows.
    for metric, start_col in (("sums_assured", 4), ("net_liability_or_current_estimate", 9)):
        metric_block(rows, ws, report_year=year, section="inforce", metric=metric, unit="HKD_million",
                     row_map=PRODUCT_ROWS_2, year_row=22, start_col=start_col,
                     linked_status="non_linked", component_type="product", source=source, checksum=checksum)
        # Sums assured uses row 32 directly as the linked subtotal. Liability/current
        # estimate uses base + additional reserve = subtotal on rows 32-34.
        controls = ((30, "non_linked", "subtotal"),
                    (32, "linked", "subtotal"),
                    (36, "total", "market_total")) if metric == "sums_assured" else (
                    (29, "non_linked", "additional_reserve"),
                    (30, "non_linked", "subtotal"),
                    (32, "linked", "base"),
                    (33, "linked", "additional_reserve"),
                    (34, "linked", "subtotal"),
                    (36, "total", "market_total"))
        for row, linked, component in controls:
            metric_block(rows, ws, report_year=year, section="inforce", metric=metric, unit="HKD_million",
                         row_map={row: "all_products"}, year_row=22, start_col=start_col,
                         linked_status=linked, component_type=component, source=source, checksum=checksum)

    # L1b: new business.
    for metric, unit, start_col in (("policy_count", "count", 4),
                                    ("office_premium", "HKD_million", 9)):
        metric_block(rows, ws, report_year=year, section="new_business", metric=metric, unit=unit,
                     row_map=PRODUCT_ROWS_NB, year_row=44, start_col=start_col,
                     linked_status="non_linked", component_type="product", source=source, checksum=checksum)
        for row, linked, component in ((51, "non_linked", "subtotal"),
                                       (53, "linked", "subtotal"),
                                       (55, "total", "market_total")):
            metric_block(rows, ws, report_year=year, section="new_business", metric=metric, unit=unit,
                         row_map={row: "all_products"}, year_row=44, start_col=start_col,
                         linked_status=linked, component_type=component, source=source, checksum=checksum)
    return rows


def main():
    rows = [row for year in (2022, 2023, 2024) for row in parse_year(year)]
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(OUTPUT)
    con.executescript("""
    DROP TABLE IF EXISTS market_amount_facts;
    CREATE TABLE market_amount_facts (
      fact_id TEXT PRIMARY KEY, report_year INTEGER NOT NULL, observation_year INTEGER NOT NULL,
      table_id TEXT NOT NULL, section TEXT NOT NULL, metric_id TEXT NOT NULL, unit TEXT NOT NULL,
      linked_status TEXT NOT NULL, insurance_type TEXT NOT NULL, component_type TEXT NOT NULL,
      value REAL, record_status TEXT NOT NULL, certification TEXT NOT NULL, schema_version TEXT NOT NULL,
      source_asset_id TEXT NOT NULL, source_file TEXT NOT NULL, source_sheet TEXT NOT NULL,
      source_locator TEXT NOT NULL, checksum_sha256 TEXT NOT NULL
    );
    CREATE INDEX idx_l1_period ON market_amount_facts(observation_year, section, metric_id);
    CREATE INDEX idx_l1_report ON market_amount_facts(report_year, linked_status, insurance_type);
    """)
    con.executemany("INSERT INTO market_amount_facts VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    con.commit()
    print("built", OUTPUT)
    print("rows", len(rows))
    print("by_year", con.execute("SELECT report_year,COUNT(*) FROM market_amount_facts GROUP BY report_year").fetchall())
    print("statuses", con.execute("SELECT record_status,COUNT(*) FROM market_amount_facts GROUP BY record_status").fetchall())
    con.close()


if __name__ == "__main__":
    main()
