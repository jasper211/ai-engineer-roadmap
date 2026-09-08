#!/usr/bin/env python3
"""Build annual L4 individual-life new-business facts for 2022-2024."""
from __future__ import annotations

import hashlib
import re
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve()
HKIA_ROOT = HERE.parents[3]
SRC = HKIA_ROOT / "12_分析框架验证_Validate_Framework/01_sources/raw/SRC-REG-IA-LTA"
DB = HERE.parents[1] / "data/annual_market_fact_layer_2022_2024.db"


def source_file(year: int) -> Path:
    return next(p for p in (SRC / str(year) / "full_annual_set").glob("Table-L*.xlsx")
                if re.match(r"Table-L4(?:_|-)", p.name))


def scalar(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "blank"
    if isinstance(value, str):
        if "N.A" in value.upper() or "不適用" in value:
            return None, "not_applicable"
        if value.strip() == "-":
            return 0.0, "reported_zero"
        return None, "unparsed"
    if isinstance(value, (int, float)):
        return float(value), "reported_zero" if float(value) == 0 else "reported"
    return None, "unparsed"


def parse(year: int):
    path = source_file(year)
    checksum = hashlib.sha256(path.read_bytes()).hexdigest()
    ws = load_workbook(path, data_only=True).active
    shift = 1 if year == 2024 else 0
    rows = []

    def emit(metric, unit, linked, participation, payment, product, component, row, col, year_row):
        observation_year = ws.cell(year_row, col).value
        if not isinstance(observation_year, int):
            raise ValueError(f"bad year {path.name}:{ws.cell(year_row, col).coordinate}={observation_year!r}")
        value, status = scalar(ws.cell(row, col).value)
        locator = ws.cell(row, col).coordinate
        fact_id = ":".join(map(str, ("L4", year, observation_year, metric, linked,
                                            participation, payment, product, component, locator)))
        rows.append((fact_id, year, observation_year, "L4", "new_business", metric, unit,
                     linked, participation, payment, product, component, value, status,
                     "certified", "rbc" if year >= 2024 else "pre_rbc", path.name,
                     path.name, ws.title, locator, checksum))

    # Non-linked: participating vs other, four product types plus subtotal.
    products = ["whole_life", "endowment", "term", "other"]
    blocks = [
        ("policy_count", "count", 7 + shift, [9 + shift, 10 + shift, 11 + shift, 12 + shift], 14 + shift),
        ("office_premium", "HKD_million", 17 + shift, [20 + shift, 21 + shift, 22 + shift, 23 + shift], 25 + shift),
    ]
    for metric, unit, year_row, detail_rows, total_row in blocks:
        for participation, start_col in (("participating", 4), ("other", 11)):
            for row, product in zip(detail_rows, products):
                for col in range(start_col, start_col + 5):
                    emit(metric, unit, "non_linked", participation, "not_applicable", product, "product", row, col, year_row)
            for col in range(start_col, start_col + 5):
                emit(metric, unit, "non_linked", participation, "not_applicable", "all_products", "subtotal", total_row, col, year_row)

    # Linked: single vs regular premium, three product types plus subtotal.
    linked_products = ["whole_life", "endowment", "other"]
    linked_blocks = [
        ("policy_count", "count", 31 + shift, [33 + shift, 34 + shift, 35 + shift], 37 + shift),
        ("office_premium", "HKD_million", 40 + shift, [43 + shift, 44 + shift, 45 + shift], 47 + shift),
    ]
    for metric, unit, year_row, detail_rows, total_row in linked_blocks:
        for payment, start_col in (("single", 4), ("regular", 11)):
            for row, product in zip(detail_rows, linked_products):
                for col in range(start_col, start_col + 5):
                    emit(metric, unit, "linked", "not_applicable", payment, product, "product", row, col, year_row)
            for col in range(start_col, start_col + 5):
                emit(metric, unit, "linked", "not_applicable", payment, "all_products", "subtotal", total_row, col, year_row)
    return rows


def main():
    facts = [fact for year in (2022, 2023, 2024) for fact in parse(year)]
    con = sqlite3.connect(DB)
    con.executescript("""
    DROP TABLE IF EXISTS new_business_detail_facts;
    CREATE TABLE new_business_detail_facts(
      fact_id TEXT PRIMARY KEY, report_year INTEGER NOT NULL, observation_year INTEGER NOT NULL,
      table_id TEXT NOT NULL, section TEXT NOT NULL, metric_id TEXT NOT NULL, unit TEXT NOT NULL,
      linked_status TEXT NOT NULL, participation_status TEXT NOT NULL, payment_basis TEXT NOT NULL,
      insurance_type TEXT NOT NULL, component_type TEXT NOT NULL, value REAL,
      record_status TEXT NOT NULL, certification TEXT NOT NULL, schema_version TEXT NOT NULL,
      source_asset_id TEXT NOT NULL, source_file TEXT NOT NULL, source_sheet TEXT NOT NULL,
      source_locator TEXT NOT NULL, checksum_sha256 TEXT NOT NULL);
    CREATE INDEX idx_l4_period ON new_business_detail_facts(observation_year, metric_id, linked_status);
    """)
    con.executemany("INSERT INTO new_business_detail_facts VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", facts)
    con.commit()
    print("rows", len(facts), "by_year", con.execute(
        "select report_year,count(*) from new_business_detail_facts group by report_year").fetchall())
    con.close()


if __name__ == "__main__":
    main()
