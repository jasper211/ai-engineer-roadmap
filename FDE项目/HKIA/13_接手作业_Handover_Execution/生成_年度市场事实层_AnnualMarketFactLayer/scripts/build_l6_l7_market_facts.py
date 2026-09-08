#!/usr/bin/env python3
"""Build L6 group/retirement and L7 annuity/other annual market facts."""
from __future__ import annotations
import hashlib, re, sqlite3
from pathlib import Path
from openpyxl import load_workbook

HERE=Path(__file__).resolve(); ROOT=HERE.parents[3]
SRC=ROOT/'12_分析框架验证_Validate_Framework/01_sources/raw/SRC-REG-IA-LTA'
DB=HERE.parents[1]/'data/annual_market_fact_layer_2022_2024.db'

def file_for(y,t):
    return next(p for p in (SRC/str(y)/'full_annual_set').glob('Table-L*.xlsx') if re.match(rf'Table-L{t}(?:_|-)',p.name))
def scalar(v):
    if v is None or (isinstance(v,str) and not v.strip()): return None,'blank'
    if isinstance(v,str):
        if 'N.A' in v.upper() or '不適用' in v or v=='#N/A': return None,'not_applicable'
        return None,'unparsed'
    if isinstance(v,(int,float)): return float(v),'reported_zero' if float(v)==0 else 'reported'
    return None,'unparsed'

def parse_l6(y):
    p=file_for(y,6); ws=load_workbook(p,data_only=True).active; sha=hashlib.sha256(p.read_bytes()).hexdigest(); out=[]
    def emit(section,metric,unit,klass,payment,component,row,col,oy):
        value,status=scalar(ws.cell(row,col).value); loc=ws.cell(row,col).coordinate
        fid=':'.join(map(str,('L6',y,oy,section,metric,klass,payment,component,loc)))
        out.append((fid,y,oy,'L6',section,metric,unit,klass,payment,component,value,status,'certified','rbc' if y>=2024 else 'pre_rbc',p.name,p.name,ws.title,loc,sha))
    group_metrics=[('policy_count','count'),('lives_count','count'),('sums_assured','HKD_million'),('office_or_revenue_premium','HKD_million'),('net_liability_or_current_estimate','HKD_million')]
    group_rows=(10,11,13,14,15) if y==2024 else (9,10,12,13,14)
    for metric_row,(metric,unit) in zip(group_rows,group_metrics):
        if y<2024:
            for oy,start in zip((y-2,y-1,y),(4,8,12)):
                for klass,col in zip(('A','C','I','total'),range(start,start+4)): emit('group_inforce',metric,unit,klass,'not_applicable','class' if klass!='total' else 'total',metric_row,col,oy)
        else:
            for oy,start in ((2021,4),(2022,8),(2023,12)):
                for klass,col in zip(('A','C','I','total'),range(start,start+4)): emit('group_inforce',metric,unit,klass,'not_applicable','class' if klass!='total' else 'total',metric_row,col,oy)
            emit('group_inforce',metric,unit,'group_life_total','not_applicable','total',metric_row,16,2024)
    if y<2024:
        rr=20; retirement_metrics=[('policy_count','count'),('lives_count','count'),('contributions','HKD_million'),('unit_liabilities','HKD_million'),('non_unit_liabilities','HKD_million'),('net_liabilities','HKD_million')]
        for row,(metric,unit) in zip((20,21,23,24,25,26),retirement_metrics):
            for oy,start in zip((y-2,y-1,y),(4,7,10)):
                for klass,col in zip(('G','H','total'),range(start,start+3)): emit('retirement_inforce',metric,unit,klass,'not_applicable','class' if klass!='total' else 'total',row,col,oy)
    else:
        retirement_metrics=[(21,'policy_count','count','not_applicable'),(22,'lives_count','count','not_applicable'),(23,'scheme_count','count','not_applicable'),(25,'contributions','HKD_million','single'),(26,'contributions','HKD_million','annual'),(27,'unit_liabilities','HKD_million','not_applicable'),(28,'non_unit_liabilities','HKD_million','not_applicable'),(29,'net_liability_or_current_estimate','HKD_million','not_applicable')]
        for row,metric,unit,payment in retirement_metrics:
            for oy,start in ((2021,4),(2022,8),(2023,11),(2024,14)):
                for klass,col in zip(('G','H','total'),range(start,start+3)): emit('retirement_inforce',metric,unit,klass,payment,'class' if klass!='total' else 'total',row,col,oy)
        for row,metric,unit,payment in [(36,'policy_count','count','single'),(37,'policy_count','count','regular'),(38,'policy_count','count','total'),(39,'lives_count','count','not_applicable'),(42,'office_premium','HKD_million','single'),(43,'office_premium','HKD_million','regular'),(44,'office_premium','HKD_million','total')]:
            for oy,start in ((2022,8),(2023,11),(2024,14)):
                for klass,col in zip(('group_life','other_group','total'),range(start,start+3)): emit('group_new_business',metric,unit,klass,payment,'segment' if klass!='total' else 'total',row,col,oy)
    return out

def parse_l7(y):
    p=file_for(y,7); ws=load_workbook(p,data_only=True).active; sha=hashlib.sha256(p.read_bytes()).hexdigest(); out=[]; shift=1 if y==2024 else 0
    def emit(section,metric,unit,linked,product,payment,component,row,col,oy):
        value,status=scalar(ws.cell(row,col).value); loc=ws.cell(row,col).coordinate
        fid=':'.join(map(str,('L7',y,oy,section,metric,linked,product,payment,component,loc)))
        out.append((fid,y,oy,'L7',section,metric,unit,linked,product,payment,component,value,status,'certified','rbc' if y>=2024 else 'pre_rbc',p.name,p.name,ws.title,loc,sha))
    yr=8+shift; years=[ws.cell(yr,c).value for c in (4,5,6)]
    products=[('individual_annuity_non_linked','non_linked',11+shift,'product'),('individual_annuity_linked','linked',12+shift,'product'),('group_annuity','not_applicable',13+shift,'product'),('individual_annuity_total','total',14+shift,'subtotal'),('permanent_health','not_applicable',17+shift,'product'),('tontines','not_applicable',18+shift,'product'),('capital_redemption','not_applicable',19+shift,'product'),('market_total','total',22+shift,'market_total')]
    for metric,unit,start in [('policy_count','count',4),('office_or_revenue_premium','HKD_million',7),('net_liability_or_current_estimate','HKD_million',10)]:
        for product,linked,row,component in products:
            for oy,col in zip(years,range(start,start+3)): emit('annuity_other_inforce',metric,unit,linked,product,'not_applicable',component,row,col,oy)
    byear=27+shift; byears=[ws.cell(byear,c).value for c in (4,5,6)]
    for metric,unit,rows in [('policy_count','count',[(30+shift,'single'),(31+shift,'regular'),(32+shift,'total')]),('office_premium','HKD_million',[(36+shift,'single'),(37+shift,'regular'),(38+shift,'total')])]:
        for linked,start in (('non_linked',4),('linked',7)):
            for row,payment in rows:
                for oy,col in zip(byears,range(start,start+3)): emit('individual_annuity_new_business',metric,unit,linked,'individual_annuity',payment,'payment_basis' if payment!='total' else 'subtotal',row,col,oy)
    return out

def main():
    l6=[r for y in (2022,2023,2024) for r in parse_l6(y)]; l7=[r for y in (2022,2023,2024) for r in parse_l7(y)]
    con=sqlite3.connect(DB); con.executescript('''
    DROP TABLE IF EXISTS group_retirement_facts;
    CREATE TABLE group_retirement_facts(fact_id TEXT PRIMARY KEY,report_year INTEGER,observation_year INTEGER,table_id TEXT,section TEXT,metric_id TEXT,unit TEXT,business_class TEXT,payment_basis TEXT,component_type TEXT,value REAL,record_status TEXT,certification TEXT,schema_version TEXT,source_asset_id TEXT,source_file TEXT,source_sheet TEXT,source_locator TEXT,checksum_sha256 TEXT);
    DROP TABLE IF EXISTS annuity_other_facts;
    CREATE TABLE annuity_other_facts(fact_id TEXT PRIMARY KEY,report_year INTEGER,observation_year INTEGER,table_id TEXT,section TEXT,metric_id TEXT,unit TEXT,linked_status TEXT,insurance_type TEXT,payment_basis TEXT,component_type TEXT,value REAL,record_status TEXT,certification TEXT,schema_version TEXT,source_asset_id TEXT,source_file TEXT,source_sheet TEXT,source_locator TEXT,checksum_sha256 TEXT);
    ''')
    con.executemany('insert into group_retirement_facts values ('+','.join('?'*19)+')',l6)
    con.executemany('insert into annuity_other_facts values ('+','.join('?'*20)+')',l7); con.commit()
    print('L6',len(l6),'L7',len(l7)); con.close()
if __name__=='__main__': main()
