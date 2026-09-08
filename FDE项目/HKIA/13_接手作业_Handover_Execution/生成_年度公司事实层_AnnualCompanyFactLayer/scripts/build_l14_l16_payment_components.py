#!/usr/bin/env python3
"""Add a lossless L14-L16 payment-basis component table without changing legacy company_facts."""
from __future__ import annotations
import hashlib, re, sqlite3
from pathlib import Path
from openpyxl import load_workbook

HERE=Path(__file__).resolve(); ROOT=HERE.parents[3]
SRC=ROOT/'12_分析框架验证_Validate_Framework/01_sources/raw/SRC-REG-IA-LTA'
DB=HERE.parents[1]/'data/annual_company_fact_layer_2022_2024.db'
SUBJECT={'L14':'non_linked_individual_life_nb','L15':'linked_individual_life_nb','L16':'total_individual_life_nb'}

def file_for(y,t): return next(p for p in (SRC/str(y)/'full_annual_set').glob('Table-L*.xlsx') if re.match(rf'Table-{t}(?:_|-)',p.name))
def scalar(v):
    if v is None: return None,'blank'
    if isinstance(v,str):
        if 'N.A' in v.upper() or '不適用' in v: return None,'not_applicable'
        if v.strip()=='-': return 0.0,'reported_zero'
        return None,'unparsed'
    return float(v),'reported_zero' if float(v)==0 else 'reported'

def parse(y,t):
    p=file_for(y,t); ws=load_workbook(p,data_only=True).active; sha=hashlib.sha256(p.read_bytes()).hexdigest(); out=[]
    if y<2024: start=10; zh,en=1,2; cols=[('policy_count','count','annual',4),('policy_count','count','single',5),('office_premium','HKD_thousand','annual',7),('office_premium','HKD_thousand','single',8)]
    else: start=11; zh,en=2,3; cols=[('policy_count','count','single',5),('policy_count','count','annual',6),('office_premium','HKD_thousand','single',8),('office_premium','HKD_thousand','annual',9)]
    for row in range(start,ws.max_row+1):
        name_en=ws.cell(row,en).value; name_zh=ws.cell(row,zh).value
        combined=f'{name_zh or ""}{name_en or ""}'
        if re.search(r'^(註|Note)',combined.strip(),re.I): break
        if not name_en and not name_zh: continue
        scope='market_total' if ('市場總額' in combined or 'market total' in combined.lower()) else 'insurer'
        name=str(name_en or name_zh).strip()
        for metric,unit,payment,col in cols:
            value,status=scalar(ws.cell(row,col).value)
            if status in ('blank','unparsed'): continue
            loc=ws.cell(row,col).coordinate; fid=':'.join(map(str,(t,y,name,metric,payment,loc)))
            out.append((fid,y,'rbc' if y==2024 else 'pre_rbc',t,SUBJECT[t],scope,name,metric,payment,value,status,unit,p.name,ws.title,loc,sha))
    return out

def main():
    rows=[r for y in (2022,2023,2024) for t in ('L14','L15','L16') for r in parse(y,t)]
    con=sqlite3.connect(DB); con.executescript('''DROP TABLE IF EXISTS company_payment_facts;
    CREATE TABLE company_payment_facts(fact_id TEXT PRIMARY KEY,report_year INTEGER,schema_version TEXT,table_id TEXT,subject TEXT,entity_scope TEXT,insurer_name_source TEXT,metric_id TEXT,payment_basis TEXT,value REAL,record_status TEXT,unit TEXT,source_file TEXT,source_sheet TEXT,source_locator TEXT,checksum_sha256 TEXT);
    CREATE INDEX idx_company_payment ON company_payment_facts(report_year,table_id,metric_id,payment_basis,entity_scope);''')
    con.executemany('insert into company_payment_facts values ('+','.join('?'*16)+')',rows); con.commit(); print('rows',len(rows)); con.close()
if __name__=='__main__': main()
